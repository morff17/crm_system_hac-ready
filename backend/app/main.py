import io, json, shutil, time, uuid
from datetime import date, datetime, timedelta
from pathlib import Path
import httpx, redis, xlrd
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Query, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse, JSONResponse
from sqlalchemy import or_, func, cast, Date
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload
from openpyxl import load_workbook
from .config import settings
from .db import Base, engine, get_db, SessionLocal
from . import models, schemas
from .auth import current_user, require_roles
from .seed import seed
from .reports import build_xlsx, build_xls, build_pdf, build_json

cache = redis.Redis.from_url(settings.redis_url, decode_responses=True)
app = FastAPI(title="RTK IT School CRM API", version="2.0.0", description="CRM API по ТЗ ИТ Школы Ростелекома")
app.add_middleware(CORSMiddleware, allow_origins=[x.strip() for x in settings.cors_origins.split(",")], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

@app.exception_handler(Exception)
async def unhandled_exception(_: Request, exc: Exception):
    if isinstance(exc, HTTPException):
        raise exc
    return JSONResponse(status_code=500, content={"code":"INTERNAL_ERROR","detail":"Внутренняя ошибка сервиса"})

@app.on_event("startup")
def startup():
    last = None
    for _ in range(30):
        try:
            Base.metadata.create_all(engine)
            with SessionLocal() as db:
                seed(db)
            Path(settings.upload_dir).mkdir(parents=True, exist_ok=True)
            return
        except Exception as e:
            last = e
            time.sleep(2)
    raise last

def invalidate_cache():
    try:
        for key in cache.scan_iter("dashboard:*"):
            cache.delete(key)
    except Exception:
        pass

def audit(db: Session, user: dict, action: str, entity_type: str, entity_id=None, details=None):
    db.add(models.AuditLog(username=user.get("preferred_username"), action=action, entity_type=entity_type, entity_id=str(entity_id) if entity_id is not None else None, details=json.dumps(details, ensure_ascii=False) if isinstance(details,(dict,list)) else details))

def user_roles(user):
    return set(user.get("roles", []))

def manager_for_user(db: Session, user: dict):
    username = user.get("preferred_username")
    sub = user.get("sub")
    return db.query(models.Manager).filter(or_(models.Manager.keycloak_username == username, models.Manager.keycloak_user_id == sub)).first()

def scope_institutions(query, db: Session, user: dict):
    if {"admin","manager"} & user_roles(user):
        return query
    manager = manager_for_user(db, user)
    if not manager:
        return query.filter(models.Institution.id == -1)
    return query.filter(models.Institution.manager_id == manager.id)

def ensure_institution_access(db: Session, user: dict, institution_id: int):
    q = scope_institutions(db.query(models.Institution), db, user).filter(models.Institution.id == institution_id)
    inst = q.first()
    if not inst:
        raise HTTPException(404, "ВУЗ не найден или недоступен")
    return inst

def parse_date_value(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    if isinstance(v, (int, float)) and 1900 <= int(v) <= 2200: return date(int(v), 12, 31)
    s = str(v).strip()
    for f in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try: return datetime.strptime(s, f).date()
        except ValueError: pass
    return None

def read_sheet(file: UploadFile):
    filename=(file.filename or "").lower()
    if filename.endswith(".xlsx"):
        wb=load_workbook(file.file, data_only=True); ws=wb.active
        return list(ws.iter_rows(values_only=True))
    if filename.endswith(".xls"):
        raw=file.file.read(); book=xlrd.open_workbook(file_contents=raw); sh=book.sheet_by_index(0)
        return [[sh.cell_value(r,c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    raise HTTPException(400,"Для импорта поддерживаются xls/xlsx")

@app.get("/api/health")
def health(): return {"status":"ok","version":"2.0.0"}

@app.post("/api/auth/login")
async def login(body: schemas.LoginRequest):
    data={"grant_type":"password","client_id":settings.keycloak_client_id,"client_secret":settings.keycloak_client_secret,"username":body.login,"password":body.password,"scope":"openid profile email"}
    async with httpx.AsyncClient(timeout=10) as client:
        r=await client.post(f"{settings.keycloak_internal_url}/realms/{settings.keycloak_realm}/protocol/openid-connect/token",data=data)
    if r.status_code!=200: raise HTTPException(401,"Неверный логин или пароль")
    return r.json()

@app.post("/api/auth/refresh")
async def refresh(body: schemas.RefreshRequest):
    data={"grant_type":"refresh_token","client_id":settings.keycloak_client_id,"client_secret":settings.keycloak_client_secret,"refresh_token":body.refresh_token}
    async with httpx.AsyncClient(timeout=10) as client:
        r=await client.post(f"{settings.keycloak_internal_url}/realms/{settings.keycloak_realm}/protocol/openid-connect/token",data=data)
    if r.status_code!=200: raise HTTPException(401,"Сессия истекла")
    return r.json()

@app.get("/api/auth/me")
def me(user=Depends(current_user), db:Session=Depends(get_db)):
    manager=manager_for_user(db,user)
    return {"username":user.get("preferred_username"),"name":user.get("name") or user.get("preferred_username"),"roles":user.get("roles",[]),"manager_id":manager.id if manager else None}


@app.get("/api/user-cache/{key}")
def get_user_cache(key:str,user=Depends(current_user)):
    safe="".join(c for c in key if c.isalnum() or c in "-_:.")[:120]
    try:
        raw=cache.get(f"usercache:{user.get('preferred_username')}:{safe}")
        return json.loads(raw) if raw else {}
    except Exception:
        return {}

@app.put("/api/user-cache/{key}")
def put_user_cache(key:str,payload:dict,user=Depends(current_user)):
    safe="".join(c for c in key if c.isalnum() or c in "-_:.")[:120]
    try: cache.setex(f"usercache:{user.get('preferred_username')}:{safe}",86400,json.dumps(payload,ensure_ascii=False))
    except Exception: pass
    return {"ok":True}

@app.delete("/api/user-cache/{key}")
def delete_user_cache(key:str,user=Depends(current_user)):
    safe="".join(c for c in key if c.isalnum() or c in "-_:.")[:120]
    try: cache.delete(f"usercache:{user.get('preferred_username')}:{safe}")
    except Exception: pass
    return {"ok":True}

@app.get("/api/institutions")
def institutions(db:Session=Depends(get_db), q:str|None=None, region:str|None=None, status:str|None=None, manager_id:int|None=None, direction_id:int|None=None, product_id:int|None=None, date_from:date|None=None, date_to:date|None=None, page:int=Query(1,ge=1), size:int=Query(20,ge=1,le=100), user=Depends(current_user)):
    query=db.query(models.Institution).options(joinedload(models.Institution.manager), joinedload(models.Institution.interactions))
    query=scope_institutions(query,db,user)
    if q: query=query.filter(or_(models.Institution.name.ilike(f"%{q}%"),models.Institution.city.ilike(f"%{q}%"),models.Institution.region.ilike(f"%{q}%")))
    if region: query=query.filter(models.Institution.region==region)
    if status: query=query.filter(models.Institution.status==status)
    if manager_id: query=query.filter(models.Institution.manager_id==manager_id)
    if direction_id or product_id:
        query=query.join(models.Interaction)
        if direction_id: query=query.filter(models.Interaction.direction_id==direction_id)
        if product_id: query=query.filter(models.Interaction.product_id==product_id)
    if date_from: query=query.filter(cast(models.Institution.created_at,Date)>=date_from)
    if date_to: query=query.filter(cast(models.Institution.created_at,Date)<=date_to)
    total=query.distinct().count(); items=query.distinct().order_by(models.Institution.name).offset((page-1)*size).limit(size).all()
    def out(i):
        return {"id":i.id,"name":i.name,"short_name":i.short_name,"full_name":i.full_name,"region":i.region,"city":i.city,"institution_type":i.institution_type,"site":i.site,"description":i.description,"status":i.status,"vendor":i.vendor,"software":i.software,"contract_number":i.contract_number,"license_signed":i.license_signed,"license_until":i.license_until,"transfer_status":i.transfer_status,"comment":i.comment,"manager":({"id":i.manager.id,"full_name":i.manager.full_name,"email":i.manager.email,"phone":i.manager.phone} if i.manager else None),"programs_count":len(i.interactions)}
    return {"items":[out(i) for i in items],"total":total,"page":page,"size":size}

@app.post("/api/institutions", dependencies=[Depends(require_roles("manager","admin"))])
def create_institution(body:schemas.InstitutionIn, db:Session=Depends(get_db), user=Depends(current_user)):
    try:
        obj=models.Institution(**body.model_dump()); db.add(obj); db.flush(); audit(db,user,"create","institution",obj.id,{"name":obj.name}); db.commit(); invalidate_cache(); return {"id":obj.id}
    except IntegrityError:
        db.rollback(); raise HTTPException(409,"ВУЗ с таким названием уже существует")

@app.get("/api/institutions/{institution_id}")
def institution_detail(institution_id:int, db:Session=Depends(get_db), user=Depends(current_user)):
    ensure_institution_access(db,user,institution_id)
    i=db.query(models.Institution).options(joinedload(models.Institution.manager),joinedload(models.Institution.contacts),joinedload(models.Institution.interactions).joinedload(models.Interaction.direction),joinedload(models.Institution.interactions).joinedload(models.Interaction.product),joinedload(models.Institution.interactions).joinedload(models.Interaction.current_stage)).filter(models.Institution.id==institution_id).first()
    return {"id":i.id,"name":i.name,"short_name":i.short_name,"full_name":i.full_name,"region":i.region,"city":i.city,"institution_type":i.institution_type,"site":i.site,"description":i.description,"status":i.status,"vendor":i.vendor,"software":i.software,"contract_number":i.contract_number,"license_signed":i.license_signed,"license_until":i.license_until,"transfer_status":i.transfer_status,"comment":i.comment,"manager":({"id":i.manager.id,"full_name":i.manager.full_name,"email":i.manager.email,"phone":i.manager.phone} if i.manager else None),"contacts":[{"id":c.id,"full_name":c.full_name,"position":c.position,"email":c.email,"phone":c.phone} for c in i.contacts],"interactions":[{"id":x.id,"workflow_id":x.workflow_id,"direction":x.direction.name if x.direction else None,"product":x.product.name if x.product else None,"status":x.status,"current_stage":{"id":x.current_stage.id,"name":x.current_stage.name,"position":x.current_stage.position} if x.current_stage else None,"students_count":x.students_count,"streams_count":x.streams_count,"applications_count":x.applications_count,"started_at":x.started_at,"ended_at":x.ended_at} for x in i.interactions]}

@app.patch("/api/institutions/{institution_id}", dependencies=[Depends(require_roles("manager","admin"))])
def patch_institution(institution_id:int, body:schemas.InstitutionPatch, db:Session=Depends(get_db), user=Depends(current_user)):
    i=db.get(models.Institution,institution_id)
    if not i: raise HTTPException(404,"ВУЗ не найден")
    for k,v in body.model_dump(exclude_unset=True).items(): setattr(i,k,v)
    audit(db,user,"update","institution",i.id,body.model_dump(exclude_unset=True,mode="json")); db.commit(); invalidate_cache(); return {"ok":True}

@app.get("/api/catalogs")
def catalogs(db:Session=Depends(get_db), user=Depends(current_user)):
    iq=scope_institutions(db.query(models.Institution),db,user)
    ids=[x[0] for x in iq.with_entities(models.Institution.id).all()]
    managers_q=db.query(models.Manager).order_by(models.Manager.full_name)
    if not ({"admin","manager"} & user_roles(user)):
        current=manager_for_user(db,user); managers_q=managers_q.filter(models.Manager.id==current.id) if current else managers_q.filter(False)
    return {"regions":[r[0] for r in iq.with_entities(models.Institution.region).filter(models.Institution.region.isnot(None)).distinct().order_by(models.Institution.region)],"statuses":[r[0] for r in iq.with_entities(models.Institution.status).distinct().order_by(models.Institution.status)],"managers":[{"id":x.id,"name":x.full_name,"keycloak_username":x.keycloak_username} for x in managers_q],"directions":[{"id":x.id,"name":x.name} for x in db.query(models.Direction).order_by(models.Direction.name)],"products":[{"id":x.id,"name":x.name,"vendor":x.vendor,"direction_id":x.direction_id} for x in db.query(models.Product).order_by(models.Product.name)]}

@app.get("/api/dashboard")
def dashboard(db:Session=Depends(get_db), user=Depends(current_user)):
    username=user.get("preferred_username","unknown"); key=f"dashboard:v2:{username}"
    try:
        cached=cache.get(key)
        if cached:return json.loads(cached)
    except Exception: pass
    iq=scope_institutions(db.query(models.Institution),db,user); ids=[x[0] for x in iq.with_entities(models.Institution.id).all()]
    xq=db.query(models.Interaction).filter(models.Interaction.institution_id.in_(ids)) if ids else db.query(models.Interaction).filter(False)
    by_status=[{"name":s,"value":c} for s,c in iq.with_entities(models.Institution.status,func.count(models.Institution.id)).group_by(models.Institution.status).all()]
    by_region=[{"name":s or "—","value":c} for s,c in iq.with_entities(models.Institution.region,func.count(models.Institution.id)).group_by(models.Institution.region).order_by(func.count(models.Institution.id).desc()).limit(8).all()]
    by_direction=[{"name":n,"value":int(v or 0)} for n,v in db.query(models.Direction.name,func.coalesce(func.sum(models.Interaction.students_count),0)).join(models.Interaction,models.Interaction.direction_id==models.Direction.id).filter(models.Interaction.institution_id.in_(ids)).group_by(models.Direction.name).order_by(func.sum(models.Interaction.students_count).desc()).all()] if ids else []
    monthly=[]
    for m in range(1,13):
        v=xq.filter(func.extract('year',models.Interaction.started_at)==datetime.utcnow().year,func.extract('month',models.Interaction.started_at)==m).with_entities(func.coalesce(func.sum(models.Interaction.students_count),0)).scalar() or 0
        monthly.append({"month":m,"value":int(v)})
    recent=xq.options(joinedload(models.Interaction.institution),joinedload(models.Interaction.current_stage)).order_by(models.Interaction.created_at.desc()).limit(6).all()
    data={"institutions":iq.count(),"programs":xq.count(),"active_contracts":iq.filter(models.Institution.contract_number.isnot(None)).count(),"students":int(xq.with_entities(func.coalesce(func.sum(models.Interaction.students_count),0)).scalar() or 0),"by_status":by_status,"by_region":by_region,"by_direction":by_direction,"monthly":monthly,"recent":[{"id":x.id,"institution_id":x.institution_id,"institution":x.institution.name,"status":x.status,"stage":x.current_stage.name if x.current_stage else None,"created_at":x.created_at} for x in recent]}
    try: cache.setex(key,30,json.dumps(data,ensure_ascii=False,default=str))
    except Exception: pass
    return data

@app.get("/api/analytics/programs")
def analytics_programs(db:Session=Depends(get_db), user=Depends(current_user)):
    iq=scope_institutions(db.query(models.Institution),db,user); ids=[x[0] for x in iq.with_entities(models.Institution.id).all()]
    rows=db.query(models.Interaction).options(joinedload(models.Interaction.institution),joinedload(models.Interaction.direction),joinedload(models.Interaction.product)).filter(models.Interaction.institution_id.in_(ids)).order_by(models.Interaction.students_count.desc()).limit(50).all() if ids else []
    return [{"id":x.id,"program":x.product.name if x.product else "—","institution":x.institution.name,"direction":x.direction.name if x.direction else "—","students":x.students_count,"status":x.status} for x in rows]

@app.get("/api/workflows")
def workflows(db:Session=Depends(get_db), user=Depends(current_user)):
    items=db.query(models.Workflow).options(joinedload(models.Workflow.stages)).order_by(models.Workflow.id).all()
    return [{"id":w.id,"name":w.name,"active":w.active,"version":w.version,"stages":[{"id":s.id,"name":s.name,"position":s.position} for s in w.stages]} for w in items]

@app.post("/api/workflows", dependencies=[Depends(require_roles("admin"))])
def create_workflow(body:schemas.WorkflowIn, db:Session=Depends(get_db), user=Depends(current_user)):
    w=models.Workflow(name=body.name,version=1); db.add(w); db.flush(); db.add_all([models.WorkflowStage(workflow_id=w.id,name=n.strip(),position=i+1) for i,n in enumerate(body.stages) if n.strip()]); audit(db,user,"create","workflow",w.id); db.commit(); return {"id":w.id}

@app.put("/api/workflows/{workflow_id}", dependencies=[Depends(require_roles("admin"))])
def update_workflow(workflow_id:int, body:schemas.WorkflowIn, db:Session=Depends(get_db), user=Depends(current_user)):
    w=db.get(models.Workflow,workflow_id)
    if not w: raise HTTPException(404,"Workflow не найден")
    used=db.query(models.Interaction).filter(models.Interaction.workflow_id==workflow_id).count()>0
    if used:
        base=body.name.strip() or w.name.split(" (v")[0]
        version=w.version+1
        name=f"{base} (v{version})"
        while db.query(models.Workflow).filter(models.Workflow.name==name).first():
            version+=1; name=f"{base} (v{version})"
        nw=models.Workflow(name=name,version=version,active=True); db.add(nw); db.flush(); db.add_all([models.WorkflowStage(workflow_id=nw.id,name=n.strip(),position=i+1) for i,n in enumerate(body.stages) if n.strip()]); w.active=False; audit(db,user,"version","workflow",nw.id,{"from":w.id}); db.commit(); return {"ok":True,"id":nw.id,"versioned":True}
    w.name=body.name; db.query(models.WorkflowStage).filter(models.WorkflowStage.workflow_id==workflow_id).delete(); db.add_all([models.WorkflowStage(workflow_id=w.id,name=n.strip(),position=i+1) for i,n in enumerate(body.stages) if n.strip()]); audit(db,user,"update","workflow",w.id); db.commit(); return {"ok":True,"id":w.id,"versioned":False}

@app.post("/api/interactions")
def create_interaction(body:schemas.InteractionCreate, db:Session=Depends(get_db), user=Depends(current_user)):
    ensure_institution_access(db,user,body.institution_id)
    first=db.query(models.WorkflowStage).filter(models.WorkflowStage.workflow_id==body.workflow_id).order_by(models.WorkflowStage.position).first()
    if not first: raise HTTPException(400,"Workflow не содержит этапов")
    x=models.Interaction(**body.model_dump(),current_stage_id=first.id,status="В процессе"); db.add(x); db.flush(); inst=db.get(models.Institution,body.institution_id); inst.status="В процессе"; audit(db,user,"create","interaction",x.id); db.commit(); db.refresh(x); invalidate_cache(); return {"id":x.id}

@app.post("/api/interactions/{interaction_id}/transition")
def transition(interaction_id:int, body:schemas.InteractionTransition, db:Session=Depends(get_db), user=Depends(current_user)):
    x=db.query(models.Interaction).options(joinedload(models.Interaction.current_stage),joinedload(models.Interaction.workflow).joinedload(models.Workflow.stages)).filter(models.Interaction.id==interaction_id).first()
    if not x: raise HTTPException(404,"Взаимодействие не найдено")
    ensure_institution_access(db,user,x.institution_id)
    stage=db.get(models.WorkflowStage,body.to_stage_id)
    if not stage or stage.workflow_id!=x.workflow_id: raise HTTPException(400,"Некорректный этап")
    cur=x.current_stage.position if x.current_stage else 0
    if stage.position not in {cur-1,cur+1,cur} and "admin" not in user_roles(user): raise HTTPException(409,"Разрешён переход только на соседний этап")
    old=x.current_stage_id; x.current_stage_id=stage.id
    maxpos=max((s.position for s in x.workflow.stages),default=stage.position)
    x.status="Завершено" if stage.position==maxpos else "В процессе"
    inst=db.get(models.Institution,x.institution_id); inst.status=x.status
    event=models.InteractionEvent(interaction_id=x.id,from_stage_id=old,to_stage_id=stage.id,comment=body.comment,author=user.get("preferred_username")); db.add(event); db.flush(); audit(db,user,"transition","interaction",x.id,{"to_stage":stage.position}); db.commit(); invalidate_cache(); return {"event_id":event.id,"stage":{"id":stage.id,"name":stage.name,"position":stage.position},"status":x.status}

@app.get("/api/interactions/{interaction_id}/history")
def interaction_history(interaction_id:int, db:Session=Depends(get_db), user=Depends(current_user)):
    x=db.get(models.Interaction,interaction_id)
    if not x: raise HTTPException(404,"Взаимодействие не найдено")
    ensure_institution_access(db,user,x.institution_id)
    events=db.query(models.InteractionEvent).options(joinedload(models.InteractionEvent.attachments)).filter(models.InteractionEvent.interaction_id==interaction_id).order_by(models.InteractionEvent.created_at.desc()).all()
    stage_ids={s.id:s for s in db.query(models.WorkflowStage).filter(models.WorkflowStage.workflow_id==x.workflow_id).all()}
    return [{"id":e.id,"from_stage":({"id":stage_ids[e.from_stage_id].id,"name":stage_ids[e.from_stage_id].name,"position":stage_ids[e.from_stage_id].position} if e.from_stage_id in stage_ids else None),"to_stage":({"id":stage_ids[e.to_stage_id].id,"name":stage_ids[e.to_stage_id].name,"position":stage_ids[e.to_stage_id].position} if e.to_stage_id in stage_ids else None),"comment":e.comment,"author":e.author,"created_at":e.created_at,"attachments":[{"id":a.id,"name":a.original_name,"size":a.size,"content_type":a.content_type} for a in e.attachments]} for e in events]

ALLOWED_EXT={"png","jpeg","jpg","pdf","zip","gzip","gz","rar","doc","docx","xls","xlsx"}
@app.post("/api/events/{event_id}/attachments")
def upload_attachment(event_id:int, file:UploadFile=File(...), db:Session=Depends(get_db), user=Depends(current_user)):
    event=db.get(models.InteractionEvent,event_id)
    if not event: raise HTTPException(404,"Событие не найдено")
    x=db.get(models.Interaction,event.interaction_id); ensure_institution_access(db,user,x.institution_id)
    ext=(file.filename.rsplit(".",1)[-1].lower() if file.filename and "." in file.filename else "")
    if ext not in ALLOWED_EXT: raise HTTPException(400,"Недопустимый формат файла")
    stored=f"{uuid.uuid4().hex}.{ext}"; path=Path(settings.upload_dir)/stored
    size=0
    with path.open("wb") as f:
        while True:
            chunk=file.file.read(1024*1024)
            if not chunk: break
            size+=len(chunk)
            if size>settings.max_upload_mb*1024*1024:
                f.close(); path.unlink(missing_ok=True); raise HTTPException(413,"Файл слишком большой")
            f.write(chunk)
    a=models.Attachment(event_id=event_id,original_name=file.filename,stored_name=stored,content_type=file.content_type,size=size); db.add(a); db.flush(); audit(db,user,"upload","attachment",a.id,{"name":a.original_name}); db.commit(); return {"id":a.id,"name":a.original_name,"size":a.size}

@app.get("/api/attachments/{attachment_id}")
def download_attachment(attachment_id:int, db:Session=Depends(get_db), user=Depends(current_user)):
    a=db.get(models.Attachment,attachment_id)
    if not a: raise HTTPException(404,"Файл не найден")
    x=db.get(models.Interaction,a.event.interaction_id); ensure_institution_access(db,user,x.institution_id)
    return FileResponse(Path(settings.upload_dir)/a.stored_name,media_type=a.content_type,filename=a.original_name)

@app.post("/api/import/{catalog}", dependencies=[Depends(require_roles("manager","admin"))])
def import_catalog(catalog:str, file:UploadFile=File(...), db:Session=Depends(get_db), user=Depends(current_user)):
    rows=read_sheet(file)
    if not rows: raise HTTPException(400,"Пустой файл")
    headers=[str(x or "").strip() for x in rows[0]]; created=updated=0
    if catalog=="institutions":
        aliases={"Название ВУЗа":"name","Наименование вуза":"name","Регион":"region","Город":"city","Вендор":"vendor","ПО":"software","Номер договора":"contract_number","Подписание лицензии":"license_signed","Срок действия лицензии (год)":"license_until","Срок действия лицензии":"license_until","Статус по передачи":"transfer_status","ФИО Менеджера":"manager_name","Ответственные от ВУЗа":"contact_name","Комментарий":"comment"}
        idx={aliases[h]:i for i,h in enumerate(headers) if h in aliases}
        if "name" not in idx: raise HTTPException(400,"Нет колонки 'Название ВУЗа'")
        for row in rows[1:]:
            name=row[idx["name"]] if idx["name"]<len(row) else None
            if not name: continue
            obj=db.query(models.Institution).filter(models.Institution.name==str(name).strip()).first()
            if not obj: obj=models.Institution(name=str(name).strip()); db.add(obj); db.flush(); created+=1
            else: updated+=1
            for key,i in idx.items():
                if i>=len(row) or row[i] in (None,"") or key in {"name","manager_name","contact_name"}: continue
                if key in {"license_signed","license_until"}: setattr(obj,key,parse_date_value(row[i]))
                else: setattr(obj,key,str(row[i]).strip())
            if "manager_name" in idx and idx["manager_name"]<len(row) and row[idx["manager_name"]]:
                n=str(row[idx["manager_name"]]).strip(); m=db.query(models.Manager).filter(models.Manager.full_name==n).first()
                if not m: m=models.Manager(full_name=n); db.add(m); db.flush()
                obj.manager_id=m.id
            if "contact_name" in idx and idx["contact_name"]<len(row) and row[idx["contact_name"]]:
                n=str(row[idx["contact_name"]]).strip()
                if not db.query(models.InstitutionContact).filter_by(institution_id=obj.id,full_name=n).first(): db.add(models.InstitutionContact(institution_id=obj.id,full_name=n))
    elif catalog=="directions":
        key=next((h for h in headers if h.lower() in {"направление","ит-направление","название"}),None)
        if not key: raise HTTPException(400,"Нужна колонка 'ИТ-направление' или 'Название'")
        i=headers.index(key)
        for row in rows[1:]:
            if i>=len(row) or not row[i]:continue
            n=str(row[i]).strip(); obj=db.query(models.Direction).filter_by(name=n).first()
            if obj: updated+=1
            else: db.add(models.Direction(name=n)); created+=1
    elif catalog=="products":
        aliases={"ИТ-продукт":"name","Продукт":"name","ПО":"name","Название":"name","Вендор":"vendor","ИТ-направление":"direction_name","Направление":"direction_name"}; idx={aliases[h]:i for i,h in enumerate(headers) if h in aliases}
        if "name" not in idx: raise HTTPException(400,"Нужна колонка продукта")
        for row in rows[1:]:
            if idx["name"]>=len(row) or not row[idx["name"]]:continue
            n=str(row[idx["name"]]).strip(); obj=db.query(models.Product).filter_by(name=n).first()
            if not obj: obj=models.Product(name=n);db.add(obj);db.flush();created+=1
            else:updated+=1
            if "vendor" in idx and idx["vendor"]<len(row) and row[idx["vendor"]]:obj.vendor=str(row[idx["vendor"]]).strip()
            if "direction_name" in idx and idx["direction_name"]<len(row) and row[idx["direction_name"]]:
                dn=str(row[idx["direction_name"]]).strip(); d=db.query(models.Direction).filter_by(name=dn).first()
                if not d:d=models.Direction(name=dn);db.add(d);db.flush()
                obj.direction_id=d.id
    elif catalog=="managers":
        aliases={"ФИО":"full_name","ФИО Менеджера":"full_name","Ответственный":"full_name","Email":"email","Телефон":"phone","Логин":"keycloak_username"};idx={aliases[h]:i for i,h in enumerate(headers) if h in aliases}
        if "full_name" not in idx: raise HTTPException(400,"Нужна колонка ФИО")
        for row in rows[1:]:
            if idx["full_name"]>=len(row) or not row[idx["full_name"]]:continue
            n=str(row[idx["full_name"]]).strip();obj=db.query(models.Manager).filter_by(full_name=n).first()
            if not obj:obj=models.Manager(full_name=n);db.add(obj);created+=1
            else:updated+=1
            for k,i in idx.items():
                if k!="full_name" and i<len(row) and row[i] not in (None,""):setattr(obj,k,str(row[i]).strip())
    else: raise HTTPException(400,"catalog: institutions, directions, products или managers")
    audit(db,user,"import",catalog,None,{"created":created,"updated":updated}); db.commit(); invalidate_cache(); return {"catalog":catalog,"created":created,"updated":updated}

# backwards-compatible endpoint
@app.post("/api/import/institutions", include_in_schema=False, dependencies=[Depends(require_roles("manager","admin"))])
def import_institutions_compat(file:UploadFile=File(...), db:Session=Depends(get_db), user=Depends(current_user)):
    return import_catalog("institutions",file,db,user)

def filtered_interactions(db,user,institution_id=None,direction_id=None,product_id=None,manager_id=None,status=None,date_from=None,date_to=None):
    q=db.query(models.Interaction).options(joinedload(models.Interaction.institution).joinedload(models.Institution.manager),joinedload(models.Interaction.direction),joinedload(models.Interaction.product)).join(models.Institution)
    q=scope_institutions(q,db,user)
    if institution_id:q=q.filter(models.Interaction.institution_id==institution_id)
    if direction_id:q=q.filter(models.Interaction.direction_id==direction_id)
    if product_id:q=q.filter(models.Interaction.product_id==product_id)
    if manager_id:q=q.filter(models.Institution.manager_id==manager_id)
    if status:q=q.filter(models.Interaction.status==status)
    if date_from:q=q.filter(or_(models.Interaction.started_at>=date_from,models.Interaction.started_at.is_(None) & (cast(models.Interaction.created_at,Date)>=date_from)))
    if date_to:q=q.filter(or_(models.Interaction.started_at<=date_to,models.Interaction.started_at.is_(None) & (cast(models.Interaction.created_at,Date)<=date_to)))
    return q.order_by(models.Interaction.created_at.desc()).all()

@app.get("/api/reports/export")
def export_report(format:str=Query("xlsx",pattern="^(xlsx|xls|pdf|json)$"),columns:str|None=None,institution_id:int|None=None,direction_id:int|None=None,product_id:int|None=None,manager_id:int|None=None,status:str|None=None,date_from:date|None=None,date_to:date|None=None,db:Session=Depends(get_db),user=Depends(current_user)):
    rows=filtered_interactions(db,user,institution_id,direction_id,product_id,manager_id,status,date_from,date_to);selected=[x.strip() for x in columns.split(",") if x.strip()] if columns else None
    builders={"xlsx":(build_xlsx,"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),"xls":(build_xls,"application/vnd.ms-excel"),"pdf":(build_pdf,"application/pdf"),"json":(build_json,"application/json")};fn,mime=builders[format];stream=fn(rows,selected);return StreamingResponse(stream,media_type=mime,headers={"Content-Disposition":f'attachment; filename="crm-report.{format}"'})

@app.post("/api/integrations/{source}/ingest", dependencies=[Depends(require_roles("manager","admin"))])
def ingest(source:str,payload:dict,db:Session=Depends(get_db),user=Depends(current_user)):
    if source not in {"lms","website"}:raise HTTPException(400,"source должен быть lms или website")
    name=payload.get("institution") or payload.get("university")
    if not name:raise HTTPException(400,"Не указано учебное заведение")
    external_id=str(payload.get("external_id")) if payload.get("external_id") is not None else None
    x=db.query(models.Interaction).filter_by(source=source,external_id=external_id).first() if external_id else None
    inst=db.query(models.Institution).filter(models.Institution.name==name).first()
    if not inst:inst=models.Institution(name=name,status="Планируется");db.add(inst);db.flush()
    inst.status="В процессе"
    direction_name=payload.get("direction");product_name=payload.get("product")
    direction=db.query(models.Direction).filter(models.Direction.name==direction_name).first() if direction_name else None
    product=db.query(models.Product).filter(models.Product.name==product_name).first() if product_name else None
    workflow=db.query(models.Workflow).filter(models.Workflow.active==True).order_by(models.Workflow.id.desc()).first();first=workflow.stages[0] if workflow and workflow.stages else None
    if not workflow:raise HTTPException(500,"Нет активного workflow")
    values={"institution_id":inst.id,"direction_id":direction.id if direction else None,"product_id":product.id if product else None,"students_count":int(payload.get("students_count",0) or 0),"streams_count":int(payload.get("streams_count",0) or 0),"applications_count":int(payload.get("applications_count",0) or 0),"external_updated_at":datetime.utcnow()}
    if x:
        for k,v in values.items():setattr(x,k,v)
        created=False
    else:
        x=models.Interaction(**values,workflow_id=workflow.id,current_stage_id=first.id if first else None,status="В процессе",source=source,external_id=external_id);db.add(x);created=True
    db.flush();audit(db,user,"ingest",source,x.id,{"external_id":external_id,"created":created});db.commit();invalidate_cache();return {"interaction_id":x.id,"source":source,"created":created}


@app.get("/api/integrations/status", dependencies=[Depends(require_roles("manager","admin"))])
def integration_status():
    return {"lms":{"configured":bool(settings.lms_api_url)},"website":{"configured":bool(settings.website_api_url)}}

@app.post("/api/integrations/{source}/sync", dependencies=[Depends(require_roles("manager","admin"))])
async def sync_integration(source:str,db:Session=Depends(get_db),user=Depends(current_user)):
    if source=="lms": url,token=settings.lms_api_url,settings.lms_api_token
    elif source=="website": url,token=settings.website_api_url,settings.website_api_token
    else: raise HTTPException(400,"source должен быть lms или website")
    if not url: raise HTTPException(409,f"URL для {source} не настроен")
    headers={"Accept":"application/json"}
    if token: headers["Authorization"]=f"Bearer {token}"
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            response=await client.get(url,headers=headers)
            response.raise_for_status()
            data=response.json()
    except (httpx.HTTPError,ValueError) as e:
        raise HTTPException(502,f"Ошибка получения данных {source}: {str(e)[:180]}")
    items=data.get("items",[]) if isinstance(data,dict) else data
    if not isinstance(items,list): raise HTTPException(502,"Интеграционный API должен вернуть массив или объект с полем items")
    result=[]
    for payload in items:
        if not isinstance(payload,dict): continue
        result.append(ingest(source,payload,db,user))
    return {"source":source,"received":len(items),"processed":len(result),"results":result[:100]}


async def keycloak_admin_token():
    if not settings.keycloak_admin_password:
        raise HTTPException(409,"KEYCLOAK_ADMIN_PASSWORD не передан backend-сервису")
    data={"grant_type":"password","client_id":"admin-cli","username":settings.keycloak_admin,"password":settings.keycloak_admin_password}
    async with httpx.AsyncClient(timeout=10) as client:
        r=await client.post(f"{settings.keycloak_internal_url}/realms/master/protocol/openid-connect/token",data=data)
    if r.status_code!=200: raise HTTPException(502,"Не удалось авторизоваться в Keycloak Admin API")
    return r.json()["access_token"]

@app.get("/api/admin/keycloak-users", dependencies=[Depends(require_roles("admin"))])
async def keycloak_users():
    token=await keycloak_admin_token(); headers={"Authorization":f"Bearer {token}"}
    async with httpx.AsyncClient(timeout=10) as client:
        r=await client.get(f"{settings.keycloak_internal_url}/admin/realms/{settings.keycloak_realm}/users?max=200",headers=headers)
        if r.status_code!=200: raise HTTPException(502,"Ошибка чтения пользователей Keycloak")
        users=r.json(); out=[]
        for u in users:
            rr=await client.get(f"{settings.keycloak_internal_url}/admin/realms/{settings.keycloak_realm}/users/{u['id']}/role-mappings/realm",headers=headers)
            roles=[x.get("name") for x in rr.json()] if rr.status_code==200 else []
            out.append({"id":u["id"],"username":u.get("username"),"firstName":u.get("firstName"),"lastName":u.get("lastName"),"enabled":u.get("enabled",True),"roles":[x for x in roles if x in {"user","manager","admin"}]})
    return out

@app.put("/api/admin/keycloak-users/{user_id}/roles", dependencies=[Depends(require_roles("admin"))])
async def keycloak_user_roles(user_id:str,body:schemas.RoleUpdate,db:Session=Depends(get_db),user=Depends(current_user)):
    wanted=set(body.roles) & {"user","manager","admin"}
    if not wanted: wanted={"user"}
    token=await keycloak_admin_token();headers={"Authorization":f"Bearer {token}","Content-Type":"application/json"}
    async with httpx.AsyncClient(timeout=10) as client:
        rr=await client.get(f"{settings.keycloak_internal_url}/admin/realms/{settings.keycloak_realm}/roles",headers=headers)
        if rr.status_code!=200: raise HTTPException(502,"Ошибка чтения ролей Keycloak")
        reps={r["name"]:r for r in rr.json() if r.get("name") in {"user","manager","admin"}}
        current=await client.get(f"{settings.keycloak_internal_url}/admin/realms/{settings.keycloak_realm}/users/{user_id}/role-mappings/realm",headers=headers)
        current_roles={r["name"]:r for r in current.json() if r.get("name") in {"user","manager","admin"}} if current.status_code==200 else {}
        to_add=[reps[r] for r in wanted-current_roles.keys() if r in reps]
        to_remove=[current_roles[r] for r in current_roles.keys()-wanted]
        if to_add:
            x=await client.post(f"{settings.keycloak_internal_url}/admin/realms/{settings.keycloak_realm}/users/{user_id}/role-mappings/realm",headers=headers,json=to_add)
            if x.status_code not in (200,204): raise HTTPException(502,"Не удалось назначить роли")
        if to_remove:
            x=await client.request("DELETE",f"{settings.keycloak_internal_url}/admin/realms/{settings.keycloak_realm}/users/{user_id}/role-mappings/realm",headers=headers,json=to_remove)
            if x.status_code not in (200,204): raise HTTPException(502,"Не удалось удалить роли")
    audit(db,user,"roles","keycloak_user",user_id,{"roles":sorted(wanted)});db.commit();return {"ok":True,"roles":sorted(wanted)}

@app.get("/api/admin/managers", dependencies=[Depends(require_roles("admin"))])
def admin_managers(db:Session=Depends(get_db)):
    return [{"id":m.id,"full_name":m.full_name,"email":m.email,"phone":m.phone,"keycloak_username":m.keycloak_username,"keycloak_user_id":m.keycloak_user_id,"institutions":db.query(models.Institution).filter_by(manager_id=m.id).count()} for m in db.query(models.Manager).order_by(models.Manager.full_name)]

@app.patch("/api/admin/managers/{manager_id}", dependencies=[Depends(require_roles("admin"))])
def admin_patch_manager(manager_id:int,body:schemas.ManagerPatch,db:Session=Depends(get_db),user=Depends(current_user)):
    m=db.get(models.Manager,manager_id)
    if not m:raise HTTPException(404,"Менеджер не найден")
    for k,v in body.model_dump(exclude_unset=True).items():setattr(m,k,v)
    audit(db,user,"update","manager",m.id,body.model_dump(exclude_unset=True));db.commit();return {"ok":True}

@app.get("/api/admin/audit", dependencies=[Depends(require_roles("admin"))])
def admin_audit(db:Session=Depends(get_db),limit:int=Query(100,ge=1,le=500)):
    rows=db.query(models.AuditLog).order_by(models.AuditLog.created_at.desc()).limit(limit).all();return [{"id":x.id,"username":x.username,"action":x.action,"entity_type":x.entity_type,"entity_id":x.entity_id,"details":x.details,"created_at":x.created_at} for x in rows]

@app.get("/api/reports/chart")
def export_chart(format:str=Query("png",pattern="^(png|pdf)$"),institution_id:int|None=None,direction_id:int|None=None,product_id:int|None=None,manager_id:int|None=None,status:str|None=None,date_from:date|None=None,date_to:date|None=None,db:Session=Depends(get_db),user=Depends(current_user)):
    rows=filtered_interactions(db,user,institution_id,direction_id,product_id,manager_id,status,date_from,date_to)
    agg={}
    for x in rows:
        name=x.direction.name if x.direction else "Без направления"
        agg[name]=agg.get(name,0)+int(x.students_count or 0)
    data=sorted(agg.items(), key=lambda x:x[1], reverse=True)
    labels=[x[0] for x in data] or ["Нет данных"]; values=[x[1] for x in data] or [0]
    if format=="pdf":
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        import os
        bio=io.BytesIO(); c=canvas.Canvas(bio,pagesize=landscape(A4)); w,h=landscape(A4); font="Helvetica"
        for p in ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf","/usr/share/fonts/dejavu/DejaVuSans.ttf"]:
            if os.path.exists(p): pdfmetrics.registerFont(TTFont("DejaVu",p));font="DejaVu";break
        c.setFont(font,15);c.drawString(40,h-45,"Студенты по ИТ-направлениям")
        maxv=max(values) or 1; y=h-90
        for label,val in zip(labels,values):
            c.setFont(font,10);c.drawString(40,y,label[:36]);c.rect(220,y-2,(w-280)*(val/maxv),12,fill=1,stroke=0);c.drawString(w-50,y,str(val));y-=28
        c.save();bio.seek(0);stream=bio;mime="application/pdf"
    else:
        from PIL import Image, ImageDraw, ImageFont
        W,H=1200,max(500,130+len(labels)*70);img=Image.new("RGB",(W,H),"white");draw=ImageDraw.Draw(img)
        font_path="/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
        try:title=ImageFont.truetype(font_path,28);body=ImageFont.truetype(font_path,20)
        except Exception:title=body=None
        draw.text((40,30),"Студенты по ИТ-направлениям",fill="black",font=title);maxv=max(values) or 1;y=100
        for label,val in zip(labels,values):
            draw.text((40,y),label[:36],fill="black",font=body);bar=int((W-430)*val/maxv);draw.rectangle((330,y,330+bar,y+28),fill=(124,58,237));draw.text((350+bar,y),str(val),fill="black",font=body);y+=65
        bio=io.BytesIO();img.save(bio,"PNG");bio.seek(0);stream=bio;mime="image/png"
    return StreamingResponse(stream,media_type=mime,headers={"Content-Disposition":f'attachment; filename="crm-chart.{format}"'})

